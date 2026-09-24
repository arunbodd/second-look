import pytest
from fastapi.testclient import TestClient

from app.main import create_app
from app.service import TriageService
from app.settings import load_settings


@pytest.fixture()
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("DB_PATH", str(tmp_path / "test.db"))
    monkeypatch.setenv("LLM_PROVIDER", "offline")
    service = TriageService(load_settings())
    with TestClient(create_app(service)) as c:
        yield c


def test_meta_and_queue(client):
    meta = client.get("/api/meta").json()
    assert meta["ai_available"] is False and meta["cases"] == 50 and meta["assessments_ready"] == 50
    assert meta["modes"]["offline"]["ready"] == 50 and meta["modes"]["ai"]["available"] is False
    q = client.get("/api/queue").json()
    assert len(q["cases"]) == 50
    b = q["briefing"]
    assert sum(l["count"] for l in b["lanes"].values()) == 50 and len(b["start_here"]) == 3
    assert q["alerts"][0]["code"] == "claim_number_collision"


def test_case_view_has_assessment_and_evidence(client):
    v = client.get("/api/cases/C1024").json()
    assert v["assessment"]["source"] == "template" and v["assessment"]["quality"]["verdict"] == "pass"
    assert v["evidence"][0]["id"] == "E1" and len(v["case"]["signals"]) == 10
    assert client.get("/api/cases/NOPE").status_code == 404


def test_decision_flow_and_audit_trail(client):
    r = client.post("/api/cases/C1024/decision", json={"decision": "accept"}, headers={"X-Investigator": "Bod"}).json()
    assert r["decision"] == "accept" and r["status"] == "in_review"
    bad = client.post("/api/cases/C1019/decision", json={"decision": "reject", "lane": "suspicious"})
    assert bad.status_code == 400  # reason required
    same = client.post("/api/cases/C1019/decision", json={"decision": "reject", "lane": "review", "reason_code": "Other (see note)"})
    assert same.status_code == 400  # same lane as the AI
    r = client.post("/api/cases/C1019/decision", json={"decision": "reject", "lane": "likely_fp", "reason_code": "Benign explanation confirmed", "note": "Family caregiver approved."}).json()
    assert r["decision"] == "reject" and r["status"] == "closed_benign" and r["final_lane"] == "likely_fp"
    assert r["notes"][0]["text"] == "Family caregiver approved."
    assert [e["type"] for e in r["timeline"]] == ["decision", "note"]
    m = client.get("/api/metrics").json()
    assert m["decided"] == 2 and m["overridden"] == 1 and m["agreement_rate"] == 0.5
    assert m["override_reasons"] == {"Benign explanation confirmed": 1}
    again = client.post("/api/cases/C1019/decision", json={"decision": "accept"})
    assert again.status_code == 400  # already decided; reset first
    r = client.post("/api/cases/C1019/decision", json={"decision": "reset"}).json()
    assert r["decision"] == "pending" and r["status"] == "new"
    m = client.get("/api/metrics").json()
    assert m["decided"] == 1 and m["agreement_rate"] == 1.0
    # needs more evidence, then the status moves on its own track
    r = client.post("/api/cases/C1026/decision", json={"decision": "needs_evidence"}).json()
    assert r["decision"] == "needs_evidence" and r["status"] == "documentation_requested"
    assert client.post("/api/cases/C1002/status", json={"status": "closed_benign"}).status_code == 400  # decide first
    r = client.post("/api/cases/C1026/status", json={"status": "escalated", "note": "SIU lead asked for it."}).json()
    assert r["status"] == "escalated" and r["decision"] == "needs_evidence"
    assert client.post("/api/cases/C1026/status", json={"status": "escalated"}).status_code == 400  # unchanged
    assert client.post("/api/cases/C1026/status", json={"status": "nope"}).status_code == 400
    r = client.post("/api/cases/C1024/status", json={"status": "closed_confirmed"}).json()
    assert r["closed"] is True and r["status_label"] == "Closed (confirmed concern)"
    q = client.get("/api/queue").json()
    row = next(c for c in q["cases"] if c["case_id"] == "C1024")
    assert row["decision"] == "accept" and row["status"] == "closed_confirmed" and row["data_gaps"]
    # deprecated action names still work
    r = client.post("/api/cases/C1033/decision", json={"action": "escalate"}).json()
    assert r["status"] == "escalated"
    md = client.get("/api/cases/C1024/export.md").text
    assert "**Decision:** Accepted AI finding" in md and "Closed (confirmed concern)" in md and "What this file cannot show" in md


def test_legacy_decision_events_still_read(client):
    """Databases written before the decision/status split keep their meaning."""
    from app.main import create_app  # noqa: F401  (client fixture already built the app)

    svc = client.app.state.service
    svc.store.add_event("C1024", "decision", {"action": "accept", "ai_lane": "suspicious", "final_lane": "suspicious", "status_after": "investigating"}, "old")
    svc.store.add_event("C1019", "decision", {"action": "override", "ai_lane": "suspicious", "final_lane": "likely_fp", "reason_code": "Benign explanation confirmed", "status_after": "closed_fp"}, "old")
    svc.store.add_event("C1026", "decision", {"action": "request_info", "ai_lane": "review", "final_lane": None, "status_after": "pending_info"}, "old")
    assert svc.case_state("C1024")["decision"] == "accept" and svc.case_state("C1024")["status"] == "in_review"
    s = svc.case_state("C1019")
    assert s["decision"] == "reject" and s["status"] == "closed_benign" and s["reason_code"] == "Benign explanation confirmed"
    assert svc.case_state("C1026")["status"] == "documentation_requested"
    svc.store.add_event("C1024", "decision", {"action": "reopen", "status_after": "open"}, "old")
    assert svc.case_state("C1024")["decision"] == "pending"


def test_excel_export(client):
    import io

    import openpyxl

    client.post("/api/cases/C1024/decision", json={"decision": "accept"})
    r = client.get("/api/export.xlsx")
    assert r.status_code == 200 and r.headers["content-type"].startswith("application/vnd.openxmlformats")
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Summary", "Queue", "Assessments", "Signals", "Score drivers", "Activity"]
    q = wb["Queue"]
    assert q.max_row == 51 and q["K2"].value.startswith("=G2*H2")
    assert wb["Summary"]["B13"].value.startswith("=SUMIFS(Queue!")
    ids = [q.cell(i, 2).value for i in range(2, 52)]
    assert "C1024" in ids and q.cell(ids.index("C1024") + 2, 13).value == "Accepted AI finding"


def test_feedback_notes_chat_and_export(client):
    r = client.post("/api/cases/C1019/feedback", json={"evidence_id": "E6", "verdict": "disagree", "comment": "approved caregiver"}).json()
    assert r["feedback"]["E6"]["verdict"] == "disagree"
    assert client.post("/api/cases/C1019/feedback", json={"evidence_id": "E999", "verdict": "agree"}).status_code == 400
    r = client.post("/api/cases/C1019/notes", json={"text": "Called the agency."}).json()
    assert r["notes"][0]["text"] == "Called the agency."
    assert client.post("/api/cases/C1019/notes", json={"text": "   "}).status_code == 400
    r = client.post("/api/cases/C1019/chat", json={"question": "why is this suspicious?"}).json()
    assert r["source"] == "offline" and r["citations"] and r["checks"]["verdict"] == "pass"
    j = client.post(f"/api/cases/C1019/chat/{r['message_id']}/judge").json()
    assert j["judge_status"] == "off" and j["verdict"] == "pass"
    md = client.get("/api/cases/C1019/export.md").text
    assert "# Case file C1019" in md and "Called the agency." in md and "why is this suspicious?" in md


def test_reset_clears_events(client):
    client.post("/api/cases/C1024/decision", json={"decision": "accept"})
    client.post("/api/reset")
    assert client.get("/api/cases/C1024").json()["state"]["status"] == "new"


def test_upload_dataset_and_return_to_sample(client, tmp_path):
    import pandas as pd

    df = pd.read_csv("data/sample_cases_synthetic.csv").head(12)
    df["case_id"] = "U" + df["case_id"].str[1:]
    csv = df.to_csv(index=False)
    r = client.post("/api/dataset?filename=my cases.csv", content=csv, headers={"Content-Type": "text/csv"})
    assert r.status_code == 200 and r.json()["rows"] == 12 and r.json()["uploaded"] is True
    meta = client.get("/api/meta").json()
    assert meta["cases"] == 12 and meta["dataset"]["name"] == "my_cases.csv"
    q = client.get("/api/queue").json()
    assert len(q["cases"]) == 12 and q["cases"][0]["case_id"].startswith("U")
    v = client.get("/api/cases/U1006").json()
    assert v["assessment"]["quality"]["verdict"] == "pass"
    assert client.get("/api/cases/C1024").status_code == 404
    bad = client.post("/api/dataset", content="case_id,foo\n1,2\n", headers={"Content-Type": "text/csv"})
    assert bad.status_code == 400 and "missing required columns" in bad.json()["detail"]
    r = client.post("/api/dataset/sample").json()
    assert r["uploaded"] is False and client.get("/api/meta").json()["cases"] == 50


def test_ai_mode_runs_beside_offline_mode(tmp_path, monkeypatch):
    """With a (mocked) model configured, both review modes are served from the same queue."""
    import json

    from app.llm.providers import OpenAICompatibleProvider, ProviderConfig
    from tests.conftest import ScriptedLLM, good_judge, good_narrative

    monkeypatch.setenv("DB_PATH", str(tmp_path / "ai.db"))
    monkeypatch.setenv("LLM_PROVIDER", "offline")
    monkeypatch.setenv("PRECOMPUTE", "false")
    service = TriageService(load_settings())
    case = service.queue.cases["C1024"]
    gen = ScriptedLLM([json.dumps(good_narrative(case))])
    judge = ScriptedLLM([good_judge()])
    service.generator = OpenAICompatibleProvider(ProviderConfig(provider="openai", model="mock", api_key="k"), gen.transport())
    service.judge = OpenAICompatibleProvider(ProviderConfig(provider="openai", model="mock-judge", api_key="k"), judge.transport())
    with TestClient(create_app(service)) as c:
        off = c.get("/api/cases/C1024?mode=offline").json()
        assert off["mode"] == "offline" and off["assessment"]["source"] == "template"
        pending = c.get("/api/cases/C1024?mode=ai").json()
        assert pending["mode"] == "ai" and pending["assessment"] is None and pending["assessment_status"] == "pending"
        ai = c.get("/api/cases/C1024?mode=ai&wait=true").json()
        assert ai["assessment"]["source"] == "llm" and ai["assessment"]["mode"] == "ai"
        assert c.get("/api/cases/C1024?mode=offline").json()["assessment"]["source"] == "template"
        q = c.get("/api/queue?mode=ai").json()
        row = next(r for r in q["cases"] if r["case_id"] == "C1024")
        assert row["source"] == "llm" and row["assessment_status"] == "ready"


def test_models_menu_and_review_run(tmp_path, monkeypatch):
    """Upload never starts the AI review; the run bar does, and models switch without a restart."""
    import json

    from app.llm import catalog
    from app.llm.providers import OpenAICompatibleProvider, ProviderConfig
    from tests.conftest import ScriptedLLM, good_judge, good_narrative

    monkeypatch.setenv("DB_PATH", str(tmp_path / "run.db"))
    monkeypatch.setenv("LLM_PROVIDER", "offline")
    monkeypatch.setenv("OPENAI_API_KEY", "k")
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    monkeypatch.delenv("GEMINI_API_KEY", raising=False)
    service = TriageService(load_settings())
    with TestClient(create_app(service)) as c:
        m = c.get("/api/models").json()
        assert {e["family"] for e in m["catalog"]} == {"openai", "anthropic", "gemini", "perplexity"}
        assert m["keys"] == {"openai": True, "anthropic": False, "gemini": False, "perplexity": False}
        assert m["current"] == {"writer": None, "judge": None}
        assert c.post("/api/review/run").status_code == 400  # no model yet
        assert c.post("/api/models", json={"writer": "claude-sonnet-5"}).status_code == 400  # no key
        assert c.post("/api/models", json={"writer": "nope"}).status_code == 400
        m = c.post("/api/models", json={"writer": "gpt-6-luna", "judge": "gpt-6-sol"}).json()
        assert m["current"] == {"writer": "gpt-6-luna", "judge": "gpt-6-sol"} and m["same_family"] is True
        assert service.settings.generator.model == "gpt-6-luna" and service.settings.judge.model == "gpt-6-sol"
        # swap the network layer for a scripted one, then run three cases
        gen = ScriptedLLM([json.dumps(good_narrative(service.queue.cases["C1024"]))])
        service.generator = OpenAICompatibleProvider(service.settings.generator, gen.transport())
        service.judge = OpenAICompatibleProvider(service.settings.judge, ScriptedLLM([good_judge()]).transport())
        st = c.get("/api/review/status").json()
        assert st["running"] is False and st["ready"] == 0
        r = c.post("/api/review/run", json={"case_ids": ["C1024", "C1019", "C1031"]}).json()
        assert r["total"] == 3
        for _ in range(200):
            st = c.get("/api/review/status").json()
            if not st["running"]:
                break
            import time
            time.sleep(0.02)
        assert st["done"] == 3 and st["ready"] == 3 and st["failed"] == 0
        assert gen.calls and gen.calls[0]["body"]["model"] == "gpt-6-luna"
        q = c.get("/api/queue?mode=ai").json()
        assert sum(1 for row in q["cases"] if row["assessment_status"] == "ready") == 3
        # running again skips what is ready
        assert c.post("/api/review/run", json={"case_ids": ["C1024"]}).json()["skipped"] == 1
        # switching models empties the AI lane; switching back restores the cache
        c.post("/api/models", json={"writer": "gpt-6-sol"})
        assert c.get("/api/review/status").json()["ready"] == 0
        c.post("/api/models", json={"writer": "gpt-6-luna"})
        assert c.get("/api/review/status").json()["ready"] == 3
        assert c.get("/api/meta").json()["models"]["writer"] == "gpt-6-luna"
        # an upload keeps the AI lane idle
        df = pd.read_csv("data/sample_cases_synthetic.csv").head(5)
        c.post("/api/dataset?filename=x.csv", content=df.to_csv(index=False), headers={"Content-Type": "text/csv"})
        st = c.get("/api/review/status").json()
        assert st["running"] is False and st["ready"] == 0 and st["cases"] == 5
        assert catalog.find("gpt-6-luna")["in"] == 0.10


import pandas as pd  # noqa: E402


def test_failed_model_calls_are_not_cached_and_can_be_retried(tmp_path, monkeypatch):
    import json
    import time

    import httpx

    from app.llm.providers import OpenAICompatibleProvider, ProviderConfig
    from tests.conftest import ScriptedLLM, good_judge, good_narrative

    monkeypatch.setenv("DB_PATH", str(tmp_path / "fail.db"))
    monkeypatch.setenv("LLM_PROVIDER", "offline")
    monkeypatch.setenv("OPENAI_API_KEY", "k")
    service = TriageService(load_settings())
    with TestClient(create_app(service)) as c:
        c.post("/api/models", json={"writer": "gpt-6-luna", "judge": "gpt-6-sol"})
        broken = httpx.MockTransport(lambda r: httpx.Response(400, json={"error": {"message": "The model `gpt-6-luna` does not exist or you do not have access to it."}}))
        service.generator = OpenAICompatibleProvider(service.settings.generator, broken)
        service.judge = OpenAICompatibleProvider(service.settings.judge, ScriptedLLM([good_judge()]).transport())
        c.post("/api/review/run", json={"case_ids": ["C1024"]})
        for _ in range(200):
            st = c.get("/api/review/status").json()
            if not st["running"]:
                break
            time.sleep(0.02)
        assert st["done"] == 1 and st["failed"] == 1 and st["ready"] == 0
        v = c.get("/api/cases/C1024?mode=ai").json()
        assert v["assessment_status"] == "failed" and v["assessment"]["source"] == "template_unavailable"
        assert "does not exist" in v["assessment"]["flags"][0]["text"]
        row = next(r for r in c.get("/api/queue?mode=ai").json()["cases"] if r["case_id"] == "C1024")
        assert row["assessment_status"] == "failed"
        assert service.cache.get_assessment("C1024", service._key(service.queue.cases["C1024"], "ai")) is None
        # fix the model and retry: the case is picked up again and cached this time
        gen = ScriptedLLM([json.dumps(good_narrative(service.queue.cases["C1024"]))])
        service.generator = OpenAICompatibleProvider(service.settings.generator, gen.transport())
        assert c.post("/api/review/run", json={"case_ids": ["C1024"]}).json()["total"] == 1
        for _ in range(200):
            st = c.get("/api/review/status").json()
            if not st["running"]:
                break
            time.sleep(0.02)
        assert st["failed"] == 0 and st["ready"] == 1
        assert c.get("/api/cases/C1024?mode=ai").json()["assessment"]["source"] == "llm"


def test_score_driver_analysis(client):
    a = client.get("/api/analysis").json()
    assert a["cases"] == 50 and len(a["signals"]) == 10 and len(a["families"]) == 5
    assert 0.5 < a["pc1_share"] < 1 and a["signal_corr_range"][0] > 0.5
    by = {s["key"]: s for s in a["signals"]}
    assert by["duplicate_service_billed"]["fired"] == 13 and by["duplicate_service_billed"]["lane_changes"] >= 0
    assert all(0 <= s["unique_share"] <= 1 for s in a["signals"])
    cols = a["matrix"]["columns"]
    assert "risk_score" not in cols and "claim_amount_usd" not in cols and len(a["matrix"]["values"]) == 10
    assert len(a["crosslogic"]) == 45
    amount = next(c for c in a["context"] if c["column"] == "claim_amount_usd")
    assert amount["value"] > 0.8
    assert next(c for c in a["context"] if c["column"] == "claim_number")["value"] == 1
    ki = a["key_indicators"]
    assert ki["assessments"] == 50 and ki["counts"]["duplicate_service_billed"] == 13
    assert len(a["scatter"]) == 50


def test_health_and_security_headers(client):
    assert client.get("/healthz").json()["ok"] is True
    ready = client.get("/readyz").json()
    assert ready["ok"] is True and ready["cases"] == 50
    r = client.get("/api/meta", headers={"X-Request-ID": "abc123"})
    h = r.headers
    assert h["x-request-id"] == "abc123" and h["x-content-type-options"] == "nosniff" and h["x-frame-options"] == "DENY"
    assert "default-src 'self'" in h["content-security-policy"] and h["cache-control"] == "no-store"
    assert client.post("/api/cases/C1024/notes", json={"text": "x" * 5000}).status_code == 422  # size limit


def test_recorded_run_replays_without_a_key(tmp_path, monkeypatch):
    """The shipped recording fills the AI review when no model is configured; stale entries read 'not run'."""
    import json

    from app.recorded import RecordedRun

    monkeypatch.setenv("DB_PATH", str(tmp_path / "rec.db"))
    monkeypatch.setenv("LLM_PROVIDER", "offline")
    base = TriageService(load_settings())
    writer = {"id": "gpt-6-luna", "label": "GPT-6 Luna", "describe": "openai:gpt-6-luna"}
    judge = {"id": "gpt-6-sol", "label": "GPT-6 Sol", "describe": "openai:gpt-6-sol"}
    probe = RecordedRun({"writer": writer, "judge": judge, "recorded_at": "2026-09-24T10:00:00+00:00"})
    cases = {}
    for cid in ("C1024", "C1015"):
        a = {**base.ensure_assessment(cid, "offline"), "source": "llm", "mode": "ai"}
        cases[cid] = {"fingerprint": probe.fingerprint(base.queue.cases[cid]), "assessment": a}
    cases["C1016"] = {"fingerprint": "stale", "assessment": cases["C1015"]["assessment"]}
    q = "Why is this case rated this way?"
    answer = {"question": q, "answer": "Recorded answer [E1].", "citations": ["E1"], "source": "llm", "model": "gpt-6-luna",
              "checks": {"verdict": "pass", "issues": []}, "warnings": [], "judge": {"verdict": "pass", "summary": "ok", "scores": {}, "issues": []}}
    path = tmp_path / "recorded.json"
    path.write_text(json.dumps({"writer": writer, "judge": judge, "recorded_at": "2026-09-24T10:00:00+00:00",
                                "cases": cases, "questions": {"C1024": [answer]}}))
    monkeypatch.setenv("RECORDED_REVIEW", str(path))
    service = TriageService(load_settings())
    with TestClient(create_app(service)) as c:
        meta = c.get("/api/meta").json()
        assert meta["ai_available"] and meta["replay"]["matched"] == 2 and "GPT-6 Luna" in meta["generator"]
        rows = {r["case_id"]: r for r in c.get("/api/queue?mode=ai").json()["cases"]}
        assert rows["C1024"]["assessment_status"] == "ready" and rows["C1016"]["assessment_status"] == "not_recorded"
        assert c.post("/api/review/run").status_code == 400  # replay never calls a model
        assert c.get("/api/models").json()["current"]["writer"] == "recorded"
        res = c.post("/api/cases/C1024/chat?mode=ai", json={"question": q}).json()
        assert res["source"] == "recorded" and "recorded" in res["model"]
        msg = [m for m in c.get("/api/cases/C1024?mode=ai").json()["state"]["chat"] if m["role"] == "assistant"][-1]
        assert msg["judge"]["verdict"] == "pass"
        other = c.post("/api/cases/C1024/chat?mode=ai", json={"question": "Who is the provider?"}).json()
        assert other["source"] == "offline"
        assert c.post("/api/models", json={"writer": "offline", "judge": "offline"}).status_code == 200
        assert c.get("/api/meta").json()["replay"] is None
        assert c.post("/api/models", json={"writer": "recorded"}).json()["current"]["writer"] == "recorded"


def test_network_guard_blocks_rebinding_and_cross_site_writes(client):
    from app import netguard

    # DNS rebinding: a foreign host name pointed at 127.0.0.1 is refused, even for reads
    assert client.get("/api/meta", headers={"host": "evil.example"}).status_code == 403
    # a web page on another site cannot trigger writes (reviews, uploads, decisions)
    r = client.post("/api/cases/C1024/notes", json={"text": "x"}, headers={"origin": "https://evil.example"})
    assert r.status_code == 403
    r = client.post("/api/cases/C1024/notes", json={"text": "x"}, headers={"sec-fetch-site": "cross-site"})
    assert r.status_code == 403
    # the app's own page and command-line clients still work
    assert client.post("/api/cases/C1024/notes", json={"text": "ok"}, headers={"origin": "http://testserver"}).status_code == 200
    assert client.post("/api/cases/C1024/notes", json={"text": "ok"}).status_code == 200
    hosts = {"localhost", "127.0.0.1", "::1"}
    assert netguard.check("GET", "127.0.0.1:8000", None, None, hosts) is None
    assert netguard.check("GET", "[::1]:8000", None, None, hosts) is None
    assert netguard.check("POST", "localhost:8000", "http://localhost:8000", "same-origin", hosts) is None
    assert netguard.check("POST", "localhost:8000", "http://localhost:3000", None, hosts)  # another local app
    assert netguard.check("GET", "192.168.1.20:8000", None, None, hosts)
    assert netguard.is_loopback("127.0.0.1") and netguard.is_loopback("::1") and not netguard.is_loopback("0.0.0.0")
